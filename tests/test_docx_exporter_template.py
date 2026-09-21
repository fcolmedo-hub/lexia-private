from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
import pytest

from services.docx_exporter import DocxExporter


def _formatted_template(path, *, marker=True):
    document = DocxDocument()
    section = document.sections[0]
    section.top_margin = Cm(4)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(4)
    section.right_margin = Cm(2)
    section.header.paragraphs[0].add_run("ESTUDIO JURÍDICO · MODELO")
    section.footer.paragraphs[0].add_run("Pie institucional")

    if marker:
        document.add_paragraph("SEÑOR JUEZ:")
        sample = document.add_paragraph()
        sample.add_run("[[CONTENIDO_LEXIA]]")
        document.add_paragraph("PROVEER DE CONFORMIDAD,")
        document.add_paragraph("SERÁ JUSTICIA.")
    else:
        sample = document.add_paragraph(
            "Texto anterior del modelo que sólo sirve para tomar la configuración de párrafo."
        )

    sample.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    sample.paragraph_format.first_line_indent = Cm(1.25)
    sample.paragraph_format.line_spacing = 1.5
    sample.paragraph_format.space_after = Pt(0)
    for run in sample.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    document.save(path)


def test_template_marker_preserves_document_and_paragraph_formatting(tmp_path):
    template = tmp_path / "modelo.docx"
    destination = tmp_path / "contestacion.docx"
    _formatted_template(template)
    original = template.read_bytes()

    DocxExporter().export_markdown_like(
        "Contestación · Audiencia",
        "# CONTESTA TRASLADO\n\nEste es el desarrollo de la contestación.",
        destination,
        template=template,
    )

    assert template.read_bytes() == original
    exported = DocxDocument(destination)
    texts = [paragraph.text for paragraph in exported.paragraphs]
    assert texts == [
        "SEÑOR JUEZ:",
        "CONTESTA TRASLADO",
        "",
        "Este es el desarrollo de la contestación.",
        "PROVEER DE CONFORMIDAD,",
        "SERÁ JUSTICIA.",
    ]
    body = next(
        paragraph for paragraph in exported.paragraphs
        if paragraph.text == "Este es el desarrollo de la contestación."
    )
    assert body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert abs(body.paragraph_format.first_line_indent - Cm(1.25)) < 1000
    assert body.paragraph_format.line_spacing == 1.5
    assert body.runs[0].font.name == "Times New Roman"
    assert body.runs[0].font.size == Pt(12)
    source_section = DocxDocument(template).sections[0]
    assert exported.sections[0].top_margin == source_section.top_margin
    assert exported.sections[0].bottom_margin == source_section.bottom_margin
    assert exported.sections[0].left_margin == source_section.left_margin
    assert exported.sections[0].right_margin == source_section.right_margin
    assert "ESTUDIO JURÍDICO · MODELO" in exported.sections[0].header.paragraphs[0].text
    assert "Pie institucional" in exported.sections[0].footer.paragraphs[0].text


def test_template_without_marker_replaces_body_but_preserves_package(tmp_path):
    template = tmp_path / "modelo_sin_marcador.docx"
    destination = tmp_path / "contestacion.docx"
    _formatted_template(template, marker=False)

    DocxExporter().export_markdown_like(
        "Contestación · Audiencia",
        "Primer párrafo de respuesta.",
        destination,
        template=template,
    )

    exported = DocxDocument(destination)
    texts = [paragraph.text for paragraph in exported.paragraphs]
    assert texts == ["Contestación · Audiencia", "Primer párrafo de respuesta."]
    assert "Texto anterior del modelo" not in "\n".join(texts)
    body = exported.paragraphs[1]
    assert body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert body.runs[0].font.name == "Times New Roman"
    assert body.runs[0].font.size == Pt(12)
    assert "ESTUDIO JURÍDICO · MODELO" in exported.sections[0].header.paragraphs[0].text
    assert "Pie institucional" in exported.sections[0].footer.paragraphs[0].text


def test_template_marker_must_occupy_its_own_paragraph(tmp_path):
    template = tmp_path / "modelo_invalido.docx"
    destination = tmp_path / "contestacion.docx"
    document = DocxDocument()
    document.add_paragraph("Texto [[CONTENIDO_LEXIA]] mezclado")
    document.save(template)

    with pytest.raises(ValueError, match="párrafo completo"):
        DocxExporter().export_markdown_like(
            "Contestación", "Contenido", destination, template=template
        )
