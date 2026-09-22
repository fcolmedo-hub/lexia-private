from zipfile import ZipFile

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

from services.docx_exporter import DocxExporter


def _export(tmp_path, content="# CONTESTA TRASLADO\n\nTexto de la contestación."):
    path = tmp_path / "contestacion.docx"
    DocxExporter().export_markdown_like(
        "Contestación · Audiencia",
        content,
        path,
    )
    return path, DocxDocument(path)


def test_export_uses_fixed_a4_mirrored_margins(tmp_path):
    path, document = _export(tmp_path)
    section = document.sections[0]

    assert abs(section.page_width - Mm(210)) < 1000
    assert abs(section.page_height - Mm(297)) < 1000
    assert abs(section.top_margin - Cm(4)) < 1000
    assert abs(section.bottom_margin - Cm(2)) < 1000
    assert abs(section.left_margin - Cm(4)) < 1000
    assert abs(section.right_margin - Cm(2)) < 1000

    with ZipFile(path) as package:
        settings = package.read("word/settings.xml").decode("utf-8")
    assert "<w:mirrorMargins" in settings


def test_export_uses_times_new_roman_12_and_26_line_pitch(tmp_path):
    _path, document = _export(tmp_path)
    normal = document.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing == Pt(25.8)
    assert normal.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY

    body = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text == "Texto de la contestación."
    )
    assert body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert body.paragraph_format.line_spacing == Pt(25.8)
    assert body.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert body.paragraph_format.space_before == Pt(0)
    assert body.paragraph_format.space_after == Pt(0)
    assert abs(body.paragraph_format.first_line_indent - Cm(1.5)) < 1000
    assert body.runs[0].font.name == "Times New Roman"
    assert body.runs[0].font.size == Pt(12)
    fonts = body.runs[0]._r.rPr.rFonts
    assert fonts.get(qn("w:ascii")) == "Times New Roman"
    assert fonts.get(qn("w:hAnsi")) == "Times New Roman"
    assert document.styles["Title"].element.pPr.find(qn("w:pBdr")) is None


def test_export_indents_only_body_paragraphs(tmp_path):
    _path, document = _export(
        tmp_path,
        "# CONTESTA TRASLADO\n\nTexto del cuerpo.\n- Elemento enumerado.",
    )
    title, heading, body, bullet = document.paragraphs

    assert title.paragraph_format.first_line_indent == 0
    assert heading.paragraph_format.first_line_indent == 0
    assert abs(body.paragraph_format.first_line_indent - Cm(1.5)) < 1000
    assert bullet.paragraph_format.first_line_indent == 0


def test_export_inserts_one_blank_line_before_each_later_chapter(tmp_path):
    _path, document = _export(
        tmp_path,
        "I - PRIMER CAPÍTULO\nÚltima oración del primer capítulo.\n"
        "II - SEGUNDO CAPÍTULO\nTexto del segundo capítulo.",
    )

    texts = [paragraph.text for paragraph in document.paragraphs]
    assert texts == [
        "CONTESTACIÓN · AUDIENCIA",
        "I - PRIMER CAPÍTULO",
        "Última oración del primer capítulo.",
        "",
        "II - SEGUNDO CAPÍTULO",
        "Texto del segundo capítulo.",
    ]
    blank_line = document.paragraphs[3]
    assert blank_line.style.name == "Normal"
    assert blank_line.paragraph_format.line_spacing == Pt(25.8)


def test_export_sets_spanish_spain_language_for_styles_and_runs(tmp_path):
    _path, document = _export(tmp_path)

    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        language = document.styles[style_name].element.rPr.find(qn("w:lang"))
        assert language is not None
        assert language.get(qn("w:val")) == "es-ES"
        assert language.get(qn("w:eastAsia")) == "es-ES"
        assert language.get(qn("w:bidi")) == "es-ES"

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            language = run._r.rPr.find(qn("w:lang"))
            assert language is not None
            assert language.get(qn("w:val")) == "es-ES"


def test_export_formats_legal_title_chapters_and_subchapters(tmp_path):
    _path, document = _export(
        tmp_path,
        "# agravios\n## Errónea fecha de ingreso\n"
        "5.2) Incorrecta valoración de la prueba\nTexto del cuerpo.",
    )
    title, chapter, generated_subchapter, explicit_subchapter, body = document.paragraphs

    assert title.text == "CONTESTACIÓN · AUDIENCIA"
    assert title.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert title.runs[0].bold is True
    assert title.runs[0].underline is True

    assert chapter.text == "I - AGRAVIOS"
    assert chapter.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert chapter.style.name == "Heading 1"
    assert chapter.runs[0].bold is True
    assert chapter.runs[0].underline is True

    assert generated_subchapter.text == "1.1) Errónea fecha de ingreso."
    assert generated_subchapter.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert generated_subchapter.style.name == "Heading 2"
    assert generated_subchapter.runs[0].bold is True
    assert generated_subchapter.runs[0].underline is not True
    assert generated_subchapter.paragraph_format.first_line_indent == 0
    assert abs(generated_subchapter.paragraph_format.left_indent - Cm(1.5)) < 1000

    assert explicit_subchapter.text == "5.2) Incorrecta valoración de la prueba."
    assert explicit_subchapter.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert explicit_subchapter.runs[0].bold is True
    assert explicit_subchapter.runs[0].underline is not True
    assert explicit_subchapter.paragraph_format.first_line_indent == 0
    assert abs(explicit_subchapter.paragraph_format.left_indent - Cm(1.5)) < 1000
    assert body.text == "Texto del cuerpo."
    assert body.paragraph_format.left_indent == 0


def test_export_sets_26_line_grid_and_footer_page_number(tmp_path):
    path, document = _export(tmp_path)
    section = document.sections[0]
    grid = section._sectPr.find(qn("w:docGrid"))
    usable_height_twips = int(Cm(29.7 - 4 - 2)) / 635
    expected_pitch = round(usable_height_twips / 26)

    assert grid is not None
    assert grid.get(qn("w:type")) == "lines"
    assert int(grid.get(qn("w:linePitch"))) == expected_pitch
    assert abs(section.footer_distance - Cm(1.2)) < 1000

    footer = section.footer.paragraphs[0]
    assert footer.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert " PAGE " in footer._p.xml
    assert footer.runs[0].font.name == "Times New Roman"
    assert footer.runs[0].font.size == Pt(12)

    with ZipFile(path) as package:
        settings = package.read("word/settings.xml").decode("utf-8")
    assert '<w:updateFields w:val="true"' in settings


def test_export_has_no_template_or_marker_requirement(tmp_path):
    path, document = _export(tmp_path, "Respuesta lista para presentar.")
    assert path.is_file()
    assert "Respuesta lista para presentar." in [
        paragraph.text for paragraph in document.paragraphs
    ]
